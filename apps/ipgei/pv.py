"""
Procès-verbal de délibération IPGEI — éditions PDF et Excel.

À la différence des documents de `documents.py`, le PV est **collectif** : il
porte toute une promotion, pas un étudiant. Il ne passe donc pas par
`DocumentOfficiel`, dont chaque ligne est rattachée à un étudiant et à un numéro
de série individuel. Il se génère à la demande, à partir des lignes de jury
telles qu'elles sont au moment du tirage.

Tant que la délibération n'est pas validée, l'édition porte la mention
« projet » : un PV non arrêté qui circulerait sans marque serait pris pour un
document définitif.
"""
import logging
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.utils import timezone

from apps.documents.services import _get_logo_b64
from core.pdf_renderer import render_pdf_bytes
from core.pdf_utils import get_institution_context

from . import services
from .models import Deliberation, LigneDeliberation

logger = logging.getLogger('siga')

TEMPLATE = 'ipgei/pv_deliberation.html'

# Paysage : sur un PV annuel, le tableau porte le rang, la moyenne, les deux
# décisions et le motif — en portrait, la colonne des motifs devient illisible.
# Le PV de semestre, plus étroit, s'en accommode aussi : garder une seule
# orientation évite deux gabarits à maintenir.
OPTIONS_PDF = {
    'page-size':    'A4',
    'orientation':  'Landscape',
    'margin-top':   '10mm',
    'margin-bottom': '10mm',
    'margin-left':  '10mm',
    'margin-right': '10mm',
    'encoding':     'UTF-8',
    'enable-local-file-access': None,
}

EN_TETES = [
    'Rang', 'Matricule', 'Nom et prénom', 'Classe', 'Moyenne générale',
    'Total coef.', 'Décision proposée', 'Décision du jury', 'Motif', 'Observations',
]


def _lignes(deliberation: Deliberation):
    return (
        deliberation.lignes
        .select_related('inscription__etudiant', 'inscription__classe')
        .prefetch_related('details__semestre')
        .order_by('rang', '-moyenne_generale')
    )


def _moyennes_par_semestre(lignes, codes: list) -> None:
    """
    Pose sur chaque ligne la moyenne de chaque semestre de l'année.

    Un PV annuel qui ne montre que la moyenne générale ne dit pas ce qui la
    compose : le jury voit un 9,80 sans savoir s'il vient d'un semestre raté
    rattrapé, ou de deux semestres également faibles — et c'est justement la
    différence qui fonde une décision.

    Calculées sur les DÉTAILS figés à la délibération, non recalculées depuis
    les notes : le PV doit dire ce que le jury a vu, même si une note a bougé
    depuis.
    """
    for ligne in lignes:
        cumuls = {code: [Decimal('0'), Decimal('0')] for code in codes}
        for detail in ligne.details.all():
            code = detail.semestre.code if detail.semestre_id else None
            if code not in cumuls or detail.note_retenue is None:
                continue
            coef = Decimal(detail.coefficient or 0)
            cumuls[code][0] += Decimal(detail.note_retenue) * coef
            cumuls[code][1] += coef
        ligne.par_semestre = [
            (cumul / poids).quantize(Decimal('0.01')) if poids else None
            for cumul, poids in (cumuls[code] for code in codes)
        ]


def _poser_rangs(deliberation: Deliberation, lignes) -> dict:
    """
    Pose sur chaque ligne son rang DANS SA CLASSE et son rang SUR LA PROMOTION.

    Un jury de semestre ne décide de rien : ce qu'il arrête, c'est un
    classement. Encore faut-il dire lequel — la promotion de MPSI est répartie
    en MPSI A et MPSI B, et le PV d'une classe ne montrait qu'un rang, sans
    dire s'il valait dans la classe ou dans la promotion. Un premier de MPSI B
    peut être quinzième de sa promotion.

    Les moyennes des étudiants DÉLIBÉRÉS sont celles du jury, telles qu'elles
    ont été figées : le PV doit dire ce que le jury a vu. Celles des autres
    classes du niveau sont calculées à la même règle — absents comptés zéro —
    pour que la comparaison ait un sens.

    Renvoie les effectifs classés, que le cartouche annonce : « 17ᵉ » ne se lit
    pas sans savoir sur combien.
    """
    from .models import InscriptionIPGEI

    semestre = deliberation.semestre
    if semestre is None:
        return {}

    figees = {l.inscription_id: l.moyenne_generale for l in lignes
              if l.moyenne_generale is not None}
    moyennes = {}
    for inscription in (InscriptionIPGEI.objects
                        .filter(actif=True,
                                annee_universitaire=deliberation.annee_universitaire,
                                classe__niveau_id=deliberation.classe.niveau_id
                                if deliberation.classe_id else None)
                        .select_related('classe')
                        if deliberation.classe_id else
                        InscriptionIPGEI.objects
                        .filter(actif=True,
                                annee_universitaire=deliberation.annee_universitaire,
                                classe__niveau__niveau=deliberation.niveau)
                        .select_related('classe')):
        valeur = figees.get(inscription.pk)
        if valeur is None:
            valeur, _, _ = services.moyenne_semestrielle(
                inscription, semestre, absents_a_zero=True)
        if valeur is not None:
            moyennes[inscription.pk] = (Decimal(valeur), inscription.classe_id)

    for ligne in lignes:
        mienne = moyennes.get(ligne.inscription_id)
        if mienne is None:
            ligne.rang_classe = ligne.rang_promotion = None
            continue
        valeur, classe_id = mienne
        ligne.rang_classe = 1 + sum(
            1 for v, c in moyennes.values() if c == classe_id and v > valeur)
        ligne.rang_promotion = 1 + sum(
            1 for v, _ in moyennes.values() if v > valeur)

    classes = {c for _, c in moyennes.values()}
    return {
        'classes_du_niveau': len(classes),
        'effectif_promotion': len(moyennes),
        'effectif_classe': (
            sum(1 for _, c in moyennes.values()
                if c == deliberation.classe_id) if deliberation.classe_id else None),
    }


def _libelle_decision(code: str) -> str:
    return dict(LigneDeliberation.DECISION_CHOICES).get(code, '')


def contexte(deliberation: Deliberation, *, base_url: str = '') -> dict:
    """Contexte de rendu du PV — partagé par le PDF et, en partie, par l'Excel."""
    lignes = list(_lignes(deliberation))
    membres = services.membres_jury_ordonnes(deliberation)

    # Colonnes de semestre : seulement sur un PV ANNUEL. Sur un PV de semestre,
    # elles répéteraient la moyenne générale dans une colonne voisine.
    codes_semestres = (
        services.codes_semestres_du_niveau(deliberation.niveau, deliberation.institution)
        if deliberation.portee == Deliberation.PORTEE_ANNUELLE else []
    )
    if codes_semestres:
        _moyennes_par_semestre(lignes, codes_semestres)

    # Un jury de SEMESTRE ne décide de rien : son PV n'a ni seuil, ni décision,
    # ni motif d'écart — trois colonnes vides sur toute la hauteur de la page.
    # Il porte à la place ce qu'il arrête vraiment : deux classements.
    semestrielle = deliberation.portee == Deliberation.PORTEE_SEMESTRE
    effectifs = _poser_rangs(deliberation, lignes) if semestrielle else {}

    # Répartition des décisions : le jury la relit avant de signer, et elle
    # évite de recompter à la main sur une promotion de 80 élèves.
    repartition = {}
    for ligne in lignes:
        if ligne.decision:
            repartition[ligne.decision] = repartition.get(ligne.decision, 0) + 1

    # L'en-tête partagé lit `inst_nom_fr`, `inst_logo_url`… : sans ce contexte
    # il se rendrait vide, sans la moindre erreur.
    donnees = get_institution_context()
    donnees.update({
        'deliberation':  deliberation,
        'institution':   deliberation.institution,
        'logo_b64':      _get_logo_b64(deliberation.institution),
        # Le titre nomme la CLASSE, comme les PV que l'institut édite depuis
        # toujours : « PV des Résultats annuels de la Classe MPSI A ». Sans
        # elle, deux PV d'un même niveau ne se distinguent qu'au contenu.
        'titre':         services.libelle_pv(deliberation),
        'codes_semestres': codes_semestres,
        'est_projet':    deliberation.statut != Deliberation.STATUT_VALIDEE,
        'niveau_libelle': 'MPSI — 1re année' if deliberation.niveau == 'MPSI' else 'MP — 2e année',
        'portee_libelle': (f'Semestre {deliberation.semestre.code}'
                           if deliberation.semestre_id else 'Année complète'),
        'lignes':        lignes,
        'effectif':      len(lignes),
        'semestrielle':  semestrielle,
        'effectif_classe':    effectifs.get('effectif_classe'),
        'effectif_promotion': effectifs.get('effectif_promotion'),
        'membres':       membres,
        'repartition':   [
            {'libelle': _libelle_decision(code), 'nombre': nombre}
            for code, nombre in sorted(repartition.items(), key=lambda c: -c[1])
        ],
        'date_edition':  date.today().strftime('%d/%m/%Y'),
        'base_url':      base_url,
    })
    return donnees


def generer_pdf(deliberation: Deliberation, *, base_url: str = '') -> bytes:
    return render_pdf_bytes(TEMPLATE, contexte(deliberation, base_url=base_url), OPTIONS_PDF)


def generer_excel(deliberation: Deliberation) -> bytes:
    """
    Classeur du PV : une feuille de résultats, une feuille de jury.

    L'Excel sert à retravailler les chiffres — il reprend donc les valeurs
    brutes (moyennes en nombres, pas en texte), là où le PDF fait foi.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    classeur = Workbook()
    feuille  = classeur.active
    feuille.title = 'Résultats'

    vert    = PatternFill('solid', fgColor='006633')
    blanc   = Font(color='FFFFFF', bold=True, size=10)
    bordure = Border(*[Side(style='thin', color='D5D5D5')] * 4)

    feuille['A1'] = (f'{services.libelle_pv(deliberation)} — '
                     f'{deliberation.annee_universitaire}')
    feuille['A1'].font = Font(bold=True, size=13, color='006633')
    # Le seuil ne sert qu'à proposer une décision : un PV de semestre n'en
    # prononce aucune, l'annoncer laisserait croire à une validation.
    semestrielle = deliberation.portee == Deliberation.PORTEE_SEMESTRE
    feuille['A2'] = (f'{"Semestre " + deliberation.semestre.code if deliberation.semestre_id else "Année complète"}'
                     + ('' if semestrielle else f' · Seuil {deliberation.seuil_validation}')
                     + f' · {deliberation.get_statut_display()}'
                     f'{" — PROJET, non validé" if deliberation.statut != Deliberation.STATUT_VALIDEE else ""}')
    feuille['A2'].font = Font(size=9, italic=True, color='555555')

    # Mêmes colonnes que le PDF, semestres compris : le classeur sert à
    # retrier et à recompter, il ne doit pas montrer moins que la feuille
    # signée.
    codes_semestres = (
        services.codes_semestres_du_niveau(deliberation.niveau, deliberation.institution)
        if deliberation.portee == Deliberation.PORTEE_ANNUELLE else []
    )
    # Sur un PV de semestre, la colonne « Rang » du jury répéterait l'une des
    # deux colonnes de rang : elle cède la place.
    tete = EN_TETES[1:4] if semestrielle else EN_TETES[:4]
    # Plus de « Décision proposée » : le procès-verbal rend ce que le jury a
    # prononcé, pas ce que le calcul suggérait. La proposition reste visible à
    # l'écran, où elle sert à délibérer ; sur la pièce signée, elle laissait
    # croire à deux décisions concurrentes.
    queue = (['Rang classe', 'Rang promotion', 'Observations'] if semestrielle
             else EN_TETES[7:])
    # Le total des coefficients ne varie pas d'un étudiant à l'autre dans une
    # même classe : sur un PV de semestre, la colonne répétait vingt-cinq fois
    # le même nombre. Le PV annuel la garde — la moyenne y agrège deux
    # semestres, et le poids permet de la refaire.
    milieu = EN_TETES[4:5] if semestrielle else EN_TETES[4:6]
    entetes = tete + list(codes_semestres) + milieu + queue
    for colonne, libelle in enumerate(entetes, start=1):
        cellule = feuille.cell(row=4, column=colonne, value=libelle)
        cellule.fill, cellule.font, cellule.border = vert, blanc, bordure
        cellule.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    lignes = list(_lignes(deliberation))
    if codes_semestres:
        _moyennes_par_semestre(lignes, codes_semestres)
    if semestrielle:
        _poser_rangs(deliberation, lignes)

    for index, ligne in enumerate(lignes, start=5):
        etudiant = ligne.inscription.etudiant
        valeurs = [
            *([] if semestrielle else [ligne.rang]),
            etudiant.matricule,
            getattr(etudiant, 'nom_display', None) or etudiant.nom,
            ligne.inscription.classe.nom,
            *[float(m) if m is not None else None
              for m in getattr(ligne, 'par_semestre', [])],
            float(ligne.moyenne_generale) if ligne.moyenne_generale is not None else None,
            *([] if semestrielle else
              [float(ligne.total_coefficients)
               if ligne.total_coefficients is not None else None]),
            # Les rangs restent des NOMBRES : le classeur sert à retrier, et
            # « 17e/25 » ne se trie pas.
            *([getattr(ligne, 'rang_classe', None),
               getattr(ligne, 'rang_promotion', None)] if semestrielle
              else [_libelle_decision(ligne.decision),
                    ligne.motif_ajustement]),
            ligne.observations,
        ]
        # Les colonnes numériques suivent le nombre de semestres insérés : les
        # figer à (5, 6) laissait la moyenne générale en texte dès qu'une
        # colonne s'ajoutait devant.
        # Le décalage suit la colonne « Rang » retirée sur un PV de semestre.
        depart = 4 if semestrielle else 5
        numeriques = set(range(
            depart, depart + len(codes_semestres) + (1 if semestrielle else 2)))
        for colonne, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(row=index, column=colonne, value=valeur)
            cellule.border = bordure
            if colonne in numeriques:
                cellule.number_format = '0.00'

    largeurs = (([16, 30, 12] if semestrielle else [7, 16, 30, 12])
                + [9] * len(codes_semestres)
                + ([15] if semestrielle else [15, 12])
                + ([14, 16, 26] if semestrielle else [24, 34, 26]))
    for index_colonne, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[
            get_column_letter(index_colonne)].width = largeur
    feuille.freeze_panes = 'A5'

    jury = classeur.create_sheet('Jury')
    for colonne, libelle in enumerate(['Rôle', 'Membre', 'Signature'], start=1):
        cellule = jury.cell(row=1, column=colonne, value=libelle)
        cellule.fill, cellule.font = vert, blanc
    for index, membre in enumerate(services.membres_jury_ordonnes(deliberation), start=2):
        jury.cell(row=index, column=1, value=membre.get_role_display())
        jury.cell(row=index, column=2,
                  value=membre.utilisateur.name or membre.utilisateur.username)
        jury.cell(row=index, column=3,
                  value=(timezone.localtime(membre.signature_le).strftime('%d/%m/%Y %H:%M')
                         if membre.a_signe else 'Non signé'))
    for colonne, largeur in zip('ABC', (28, 34, 20)):
        jury.column_dimensions[colonne].width = largeur

    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
