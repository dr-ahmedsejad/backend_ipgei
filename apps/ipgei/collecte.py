"""
Fiches de collecte de notes — édition PDF et Excel.

Une fiche de collecte est la feuille que l'enseignant emporte en salle : la
liste de ses étudiants, une colonne de notes vide et une colonne d'observations.
Il la remplit à la main, la signe, et la scolarité reporte ensuite les notes à
l'écran. C'est le pendant papier de la grille de saisie.

Trois exigences en découlent, qui commandent tout ce module :

  * **Même ensemble et même ordre que la grille de saisie.** L'enseignant lit
    sa feuille ligne à ligne pendant qu'un agent tape à l'écran. Un ordre
    différent entre les deux, et les notes se décalent sans que rien ne
    l'annonce. D'où `LIGNES_ORDRE`, partagé avec `views.NoteViewSet.grille`.
  * **Une fiche par matière.** Un enseignant corrige une matière, pas un
    semestre : mêler deux matières sur une feuille l'oblige à trier.
  * **L'anonymat est un mode, pas une variante.** Sous anonymat la fiche ne
    porte QUE le numéro — ni matricule ni nom, sinon elle ne protège rien.

Repris du dispositif SIGA (`evaluations/services/note_reports.py`), transposé
aux objets IPGEI : `Matiere` au lieu de `EM`, `Classe` au lieu de filière, et
`InscriptionIPGEI` au lieu de l'inscription pédagogique LMD.
"""
import io
import logging
import os
from datetime import date

from django.core.exceptions import ValidationError

from core.pdf_renderer import render_pdf_bytes
from core.pdf_utils import get_institution_context

from . import services
from .models import (
    AnonymatIPGEI, InscriptionIPGEI, Matiere, Note, ParametresIPGEI,
    SessionEvaluationIPGEI,
)

logger = logging.getLogger('siga')

TEMPLATE = 'ipgei/fiche_collecte.html'

#: Ordre de la grille de saisie. Toute divergence décale le report des notes.
LIGNES_ORDRE = ('numero_ordre', 'etudiant__nom')

TYPE_DS   = 'DS'
TYPE_TP   = 'TP'
TYPE_EXAM = 'EXAM'
#: Le rattrapage est un TYPE D'ÉPREUVE et non une campagne à choisir : c'est
#: ainsi que l'écran le présente. La campagne se déduit du type — les trois
#: premiers relèvent de la session normale, celui-ci de la seconde session.
TYPE_RATT = 'RATT'
TYPES_NOTE = (TYPE_DS, TYPE_TP, TYPE_EXAM, TYPE_RATT)

LIBELLES_TYPE = {
    TYPE_DS:   'Devoir surveillé',
    TYPE_TP:   'Travaux pratiques',
    TYPE_EXAM: 'Examen final',
    TYPE_RATT: 'Examen de rattrapage',
}

OPTIONS_PDF = {
    'page-size':    'A4',
    'orientation':  'Portrait',
    'margin-top':   '12mm',
    'margin-bottom': '12mm',
    'margin-left':  '14mm',
    'margin-right': '14mm',
    'encoding':     'UTF-8',
    'enable-local-file-access': None,
}


# ═════════════════════════════════════════════════════════════════════════════
# Sélection des lignes
# ═════════════════════════════════════════════════════════════════════════════
def libelle_type_note(type_note: str) -> str:
    return LIBELLES_TYPE.get(type_note, type_note)


def campagne_du_type(type_note: str) -> str:
    """Campagne de saisie dont relève une épreuve — déduite, jamais demandée."""
    return (SessionEvaluationIPGEI.TYPE_RATTRAPAGE if type_note == TYPE_RATT
            else SessionEvaluationIPGEI.TYPE_NORMALE)


def _inscriptions(classe_id: int):
    return (
        InscriptionIPGEI.objects
        .filter(classe_id=classe_id, actif=True)
        .select_related('etudiant')
        .order_by(*LIGNES_ORDRE)
    )


def _inscriptions_a_rattraper(classe_id: int, matiere: Matiere, semestre):
    """
    En seconde session, seuls les étudiants qui doivent la matière.

    Convoquer toute la classe à un rattrapage ferait signer des présents qui
    n'avaient rien à repasser, et la feuille ne dirait plus qui était attendu.
    Le seuil est celui du paramétrage courant : la fiche s'imprime AVANT le
    jury, il n'existe donc pas encore de seuil figé par une délibération.
    """
    parametres = ParametresIPGEI.pour(matiere.institution)
    seuil = parametres.seuil_validation_semestre or parametres.seuil_validation
    concernes = set(
        Note.objects
        .filter(matiere=matiere, semestre=semestre,
                inscription__classe_id=classe_id,
                moyenne__lt=seuil)
        .values_list('inscription_id', flat=True)
    )
    return [i for i in _inscriptions(classe_id) if i.pk in concernes], seuil


def lignes(semestre, classe_id: int, matiere: Matiere,
           type_note: str, anonymat: bool) -> tuple:
    """
    Lignes de la fiche et seuil retenu, dans l'ordre de la grille de saisie.

    Sous anonymat, une inscription sans numéro tiré est ÉCARTÉE plutôt que
    laissée avec une case vide : une ligne anonyme sans numéro ne peut être
    reportée sur personne.
    """
    rattrapage = type_note == TYPE_RATT
    seuil = None
    if rattrapage:
        inscriptions, seuil = _inscriptions_a_rattraper(
            classe_id, matiere, semestre)
    else:
        inscriptions = list(_inscriptions(classe_id))

    numeros = {}
    if anonymat:
        # Les numéros sont tirés par campagne : celle de l'épreuve, pas une
        # autre — un numéro de session normale ne désigne pas le même candidat
        # qu'un numéro de rattrapage.
        session = services.session_saisie(semestre, campagne_du_type(type_note))
        numeros = dict(
            AnonymatIPGEI.objects.filter(session=session)
            .values_list('inscription_id', 'numero')
        )
        if not numeros:
            raise ValidationError(
                "Aucun anonymat n'a été tiré pour cette campagne : générez-les "
                "avant d'éditer une fiche anonyme."
            )

    sortie = []
    for inscription in inscriptions:
        if anonymat:
            numero = numeros.get(inscription.pk)
            if numero is None:
                continue
            sortie.append({'anonymat': numero})
        else:
            sortie.append({
                'matricule': inscription.etudiant.matricule,
                'nom':       inscription.etudiant.nom_display,
            })
    if anonymat:
        # Trié par numéro : sous anonymat, c'est le seul repère du correcteur.
        sortie.sort(key=lambda ligne: ligne['anonymat'])
    return sortie, seuil


def matieres_de(classe, semestre, type_note: str):
    """
    Matières du semestre pour le niveau de la classe.

    Une fiche TP n'a de sens que pour une matière qui en comporte : l'éditer
    ailleurs produirait une feuille que personne ne remplira.
    """
    requete = Matiere.objects.filter(
        institution=classe.institution,
        code_semestre=semestre.code,
        actif=True,
    )
    niveau = getattr(classe, 'niveau_code', None)
    if niveau:
        requete = requete.filter(niveau_ref__code=niveau)
    if type_note == TYPE_TP:
        requete = requete.filter(has_tp=True)
    return list(requete.order_by('ordre', 'code'))


# ═════════════════════════════════════════════════════════════════════════════
# Contexte et rendu
# ═════════════════════════════════════════════════════════════════════════════
def _fiche(semestre, classe, matiere, type_note, anonymat) -> dict:
    if type_note not in TYPES_NOTE:
        raise ValidationError(f'Type de note inconnu : {type_note}.')
    if type_note == TYPE_TP and not matiere.has_tp:
        raise ValidationError(
            f'{matiere.code} ne comporte pas de TP notés : aucune fiche de TP '
            "n'a de sens pour cette matière."
        )
    contenu, seuil = lignes(semestre, classe.pk, matiere, type_note, anonymat)
    return {
        'matiere':  matiere,
        'lignes':   contenu,
        'seuil':    seuil,
        'effectif': len(contenu),
    }


def contexte(semestre, classe, matieres: list,
             type_note: str, anonymat: bool) -> dict:
    """Contexte de rendu — une fiche par matière, dans l'ordre de la maquette."""
    rattrapage = type_note == TYPE_RATT
    donnees = get_institution_context()
    donnees.update({
        'fiches':        [_fiche(semestre, classe, m, type_note, anonymat)
                          for m in matieres],
        'classe':        classe,
        'classe_label':  services.libelle_classe(classe),
        'semestre':      semestre,
        'annee_universitaire': semestre.annee_universitaire,
        'type_note':     type_note,
        'type_note_label': libelle_type_note(type_note),
        'anonymat':      anonymat,
        'rattrapage':    rattrapage,
        'date_edition':  date.today().strftime('%d/%m/%Y'),
    })
    return donnees


def generer_pdf(semestre, classe, matieres: list, type_note: str,
                anonymat: bool) -> bytes:
    return render_pdf_bytes(TEMPLATE,
                            contexte(semestre, classe, matieres, type_note, anonymat),
                            dict(OPTIONS_PDF))


#: Caractères qu'Excel refuse dans un nom d'onglet.
_INTERDITS_ONGLET = ':\\/?*[]'


def _onglet(matiere) -> str:
    """
    Nom d'onglet : l'intitulé de la matière, comme sur la fiche.

    Excel plafonne à 31 caractères, refuse `: \\ / ? * [ ]` et veut des noms
    uniques — d'où le nettoyage et la troncature. Le code servait de nom court
    jusqu'ici ; il ne parle qu'à la scolarité, et l'onglet se lit par
    l'enseignant.
    """
    libelle = ''.join(c for c in (matiere.intitule or '')
                      if c not in _INTERDITS_ONGLET).strip()
    return (libelle or matiere.code or 'Matière')[:28]


def _chemin_logo(donnees: dict) -> str:
    """
    Chemin disque du logo, tiré de l'URL `file:///` que le contexte prépare
    pour wkhtmltopdf. Excel, lui, veut un fichier à ouvrir.
    """
    url = donnees.get('inst_logo_url') or donnees.get('image_url') or ''
    chemin = url[8:] if url.startswith('file:///') else url
    return chemin if chemin and os.path.exists(chemin) else ''


#: Côté du logo dans le classeur, en pixels — celui du bandeau PDF.
_COTE_LOGO = 70

#: Un pixel vaut 9 525 EMU (English Metric Units), l'unité des dessins Office.
_EMU_PAR_PIXEL = 9525


def _emu(pixels: float) -> int:
    return int(round(pixels * _EMU_PAR_PIXEL))


def _px_colonne(largeur: float) -> float:
    """Largeur d'une colonne en pixels — Excel la compte en caractères."""
    return largeur * 7 + 5


def _px_ligne(hauteur: float) -> float:
    """Hauteur d'une ligne en pixels — Excel la compte en points (96/72)."""
    return hauteur * 4 / 3


def _logo_reduit(donnees: dict) -> bytes:
    """
    Logo ramené à la taille où il s'affiche, une fois pour tout le classeur.

    Excel embarque l'image dans CHAQUE feuille : le fichier d'origine, en pleine
    résolution, pesait près de 90 ko par matière — un demi-mégaoctet pour six
    fiches, à envoyer par courriel. Réduit, il en fait quelques-uns.
    Renvoie des octets vides si le logo est absent ou illisible.
    """
    chemin = _chemin_logo(donnees)
    if not chemin:
        return b''
    try:
        from PIL import Image as ImagePIL

        with ImagePIL.open(chemin) as image:
            image = image.convert('RGBA')
            # Deux fois la taille d'affichage : net à l'impression, léger.
            image.thumbnail((_COTE_LOGO * 2, _COTE_LOGO * 2))
            tampon = io.BytesIO()
            image.save(tampon, format='PNG')
            return tampon.getvalue()
    except Exception:                                        # noqa: BLE001
        # Un logo illisible ne doit pas emporter la fiche : le classeur sort
        # sans image plutôt que pas du tout.
        logger.exception('Logo illisible pour le classeur de collecte')
        return b''


def _entete_excel(feuille, donnees: dict, nb_colonnes: int, logo: bytes) -> int:
    """
    Bandeau d'institution — le même que sur le PDF : intitulé français à
    gauche, logo au centre, arabe à droite, filet en dessous.

    Le classeur circule autant que la feuille imprimée : sans en-tête, on ne
    savait plus de quel établissement venait le fichier reçu par courriel.
    Renvoie le numéro de la première ligne libre.
    """
    from openpyxl.drawing.image import Image
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    # Le logo occupe l'AVANT-DERNIÈRE colonne : le français prend tout ce qui
    # est à sa gauche, l'arabe tout ce qui est à sa droite. Serré dans la seule
    # colonne A — large de seize caractères, celle des matricules — l'intitulé
    # français se coupait en cinq lignes et débordait de la bande.
    colonne_logo = max(2, nb_colonnes - 1)
    lettre_logo  = get_column_letter(colonne_logo)
    derniere     = get_column_letter(nb_colonnes)

    if colonne_logo > 2:
        feuille.merge_cells(f'A1:{get_column_letter(colonne_logo - 1)}3')
    else:
        feuille.merge_cells('A1:A3')
    feuille.merge_cells(f'{lettre_logo}1:{lettre_logo}3')
    if nb_colonnes > colonne_logo:
        feuille.merge_cells(
            f'{get_column_letter(colonne_logo + 1)}1:{derniere}3')

    # Largeurs minimales du bandeau, sans jamais rétrécir celles du tableau :
    # de quoi loger le nom de l'établissement et les 70 px du logo.
    def au_moins(colonne: int, largeur: float) -> None:
        lettre = get_column_letter(colonne)
        actuelle = feuille.column_dimensions[lettre].width or 0
        if actuelle < largeur:
            feuille.column_dimensions[lettre].width = largeur

    au_moins(1, 24 if colonne_logo == 2 else 16)
    au_moins(colonne_logo, 13)
    if nb_colonnes > colonne_logo:
        au_moins(colonne_logo + 1, 26)

    # Blocs gauche et droite mis à ÉGALE largeur, faute de quoi la colonne du
    # logo n'est pas au milieu de la bande : centré dans sa propre colonne, il
    # tombait tout de même à droite du centre, et cela se voyait. On ne fait
    # qu'élargir le plus étroit — jamais rétrécir une colonne du tableau.
    # Comparaison en PIXELS et non en caractères : chaque colonne coûte cinq
    # pixels de marge, si bien que deux colonnes de cinquante caractères sont
    # plus larges qu'une seule — l'écart suffisait à décaler le logo.
    def largeur_px(colonnes) -> float:
        return sum(_px_colonne(
            feuille.column_dimensions[get_column_letter(c)].width or 0)
            for c in colonnes)

    colonnes_fr = list(range(1, colonne_logo))
    colonnes_ar = list(range(colonne_logo + 1, nb_colonnes + 1))
    if colonnes_fr and colonnes_ar:
        ecart = largeur_px(colonnes_fr) - largeur_px(colonnes_ar)
        if ecart > 0:
            au_moins(nb_colonnes,
                     (feuille.column_dimensions[derniere].width or 0) + ecart / 7)
        elif ecart < 0:
            au_moins(1, (feuille.column_dimensions['A'].width or 0) - ecart / 7)

    # Hauteur de la bande : trois lignes de texte à 11 pt, et les 70 px du logo
    # (≈ 53 pt). Une cellule fusionnée ne s'ajuste pas d'elle-même.
    for ligne in (1, 2, 3):
        feuille.row_dimensions[ligne].height = 24

    gauche = '\n'.join(filter(None, [donnees.get('inst_groupe_fr'),
                                     donnees.get('inst_nom_fr')]))
    droite = '\n'.join(filter(None, [donnees.get('inst_groupe_ar'),
                                     donnees.get('inst_nom_ar')]))
    cellule = feuille['A1']
    cellule.value = gauche
    # Même rapport que le bandeau du PDF, où l'arabe est un cran plus grand.
    cellule.font = Font(bold=True, size=11)
    cellule.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

    if nb_colonnes > colonne_logo:
        cellule = feuille.cell(1, colonne_logo + 1)
        cellule.value = droite
        cellule.font = Font(bold=True, size=12)
        cellule.alignment = Alignment(horizontal='right', vertical='center',
                                      wrap_text=True)

    if logo:
        image = Image(io.BytesIO(logo))
        image.width = image.height = _COTE_LOGO
        # Ancré au coin de la cellule, le logo se collait à gauche et en haut
        # de sa colonne : il paraissait décalé face au bandeau du PDF, où il
        # est centré. Excel ne centre pas une image, il l'accroche — on calcule
        # donc le décalage à partir de la largeur et de la hauteur réelles.
        largeur_case = _px_colonne(
            feuille.column_dimensions[lettre_logo].width or 13)
        hauteur_case = sum(_px_ligne(feuille.row_dimensions[ligne].height or 22)
                           for ligne in (1, 2, 3))
        image.anchor = OneCellAnchor(
            _from=AnchorMarker(
                col=colonne_logo - 1,
                colOff=_emu(max(0, (largeur_case - _COTE_LOGO) / 2)),
                row=0,
                rowOff=_emu(max(0, (hauteur_case - _COTE_LOGO) / 2)),
            ),
            ext=XDRPositiveSize2D(_emu(_COTE_LOGO), _emu(_COTE_LOGO)),
        )
        feuille.add_image(image)

    # Le filet du PDF, rendu par une bordure basse sur toute la largeur.
    filet = Side(border_style='thin', color='000000')
    for colonne in range(1, nb_colonnes + 1):
        feuille.cell(4, colonne).border = Border(bottom=filet)
    feuille.row_dimensions[4].height = 6
    return 5


def generer_excel(semestre, classe, matieres: list, type_note: str,
                  anonymat: bool) -> bytes:
    """
    Classeur de collecte : une feuille par matière, colonnes de notes vides.

    Même contenu et même en-tête que le PDF — c'est le support qui change, pas
    la fiche. Un enseignant qui saisit sur ordinateur renvoie le fichier ; celui
    qui corrige sur copie imprime le PDF.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    donnees = contexte(semestre, classe, matieres, type_note, anonymat)

    VERT = '006633'
    # Gris franc et non presque blanc : la feuille se remplit au stylo, les
    # lignes doivent se voir à l'impression comme à l'écran. C'est le ton du
    # tableau du PDF.
    trait = Side(border_style='thin', color='999999')
    bordure = Border(left=trait, right=trait, top=trait, bottom=trait)
    centre = Alignment(horizontal='center', vertical='center')

    classeur = openpyxl.Workbook()
    classeur.remove(classeur.active)
    logo = _logo_reduit(donnees)

    for fiche in donnees['fiches']:
        matiere = fiche['matiere']
        base = _onglet(matiere)
        nom, suite = base, 1
        while nom in classeur.sheetnames:
            suite += 1
            nom = f'{base}_{suite}'
        feuille = classeur.create_sheet(nom)

        entetes = (['N° anonymat', 'Note /20', 'Observations'] if anonymat
                   else ['Matricule', 'Nom et prénom', 'Note /20', 'Observations'])
        nb_colonnes = len(entetes)

        # Les largeurs du tableau AVANT le bandeau : celui-ci ne fait que les
        # relever si le nom de l'établissement ou le logo n'y tiennent pas.
        largeurs = [16, 12, 34] if anonymat else [16, 34, 12, 34]
        for index, largeur in enumerate(largeurs, 1):
            feuille.column_dimensions[get_column_letter(index)].width = largeur

        rang = _entete_excel(feuille, donnees, nb_colonnes, logo)

        # Titre puis bloc d'identification, dans l'ordre de la fiche papier.
        titre = feuille.cell(rang, 1, f'Fiche de collecte — {donnees["type_note_label"]}')
        titre.font = Font(bold=True, size=12, color=VERT)
        titre.alignment = Alignment(horizontal='center')
        feuille.merge_cells(start_row=rang, start_column=1,
                            end_row=rang, end_column=nb_colonnes)
        rang += 2

        for etiquette, valeur in (
            ('Classe',         donnees['classe_label']),
            ('Matière',        matiere.intitule),
            ('Semestre',       donnees['semestre'].code),
            ('Effectif',       f'{fiche["effectif"]} étudiant'
                               f'{"s" if fiche["effectif"] > 1 else ""}'),
            ('Année',          donnees['annee_universitaire']),
            ("Date d'édition", donnees['date_edition']),
        ):
            cellule = feuille.cell(rang, 1, f'{etiquette} :')
            cellule.font = Font(bold=True)
            feuille.cell(rang, 2, valeur)
            feuille.merge_cells(start_row=rang, start_column=2,
                                end_row=rang, end_column=nb_colonnes)
            rang += 1

        ligne_entete = rang + 1
        for colonne, entete in enumerate(entetes, 1):
            cellule = feuille.cell(ligne_entete, colonne, entete)
            cellule.font = Font(bold=True, color='FFFFFF')
            cellule.fill = PatternFill('solid', fgColor=VERT)
            cellule.alignment = centre
            cellule.border = bordure

        rang = ligne_entete + 1
        for ligne in fiche['lignes']:
            valeurs = ([ligne['anonymat'], None, None] if anonymat
                       else [ligne['matricule'], ligne['nom'], None, None])
            for colonne, valeur in enumerate(valeurs, 1):
                cellule = feuille.cell(rang, colonne, valeur)
                cellule.border = bordure
            rang += 1

        if fiche['seuil'] is not None:
            note = feuille.cell(
                rang + 1, 1,
                'Liste restreinte aux étudiants dont la moyenne de la matière '
                f'est inférieure à {fiche["seuil"]} / 20.')
            note.font = Font(italic=True, size=9, color='666666')
            rang += 1

        # Le bloc de signature de la fiche papier : le classeur s'imprime aussi.
        rang += 2
        signature = feuille.cell(rang, 1, "Signature de l'enseignant")
        signature.font = Font(bold=True, size=9, color='555555')
        feuille.cell(rang + 2, 1, 'Date : _____ / _____ / _______').font = \
            Font(size=9, color='555555')

        # Les en-têtes restent visibles quand la classe dépasse un écran.
        feuille.freeze_panes = feuille.cell(ligne_entete + 1, 1)

    if not classeur.sheetnames:
        classeur.create_sheet('Vide')

    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
